#!/usr/bin/env python3
"""版面自检：用真实字体度量估算每个文本框的换行与高度，找出溢出与越界。"""
import sys
from pptx import Presentation
from pptx.util import Emu
from PIL import ImageFont

PING = ("/System/Library/AssetsV2/com_apple_MobileAsset_Font8/"
        "86ba2c91f017a3749571a82f2c6d890ac7ffb2fb.asset/AssetData/PingFang.ttc")
SONG = "/System/Library/Fonts/Supplemental/Songti.ttc"
MONO = "/System/Library/Fonts/SFNSMono.ttf"

_CACHE = {}


def face(font_name, bold, size):
    key = (font_name, bold, size)
    if key in _CACHE:
        return _CACHE[key]
    if font_name == "Songti SC":
        path, idx = SONG, 0
    elif font_name == "SF Mono":
        path, idx = MONO, 0
    else:
        path, idx = PING, (11 if bold else 3)
    f = ImageFont.truetype(path, max(1, int(round(size))), index=idx)
    _CACHE[key] = f
    return f


def wrap_lines(text, font, width_pt):
    """按字符贪心折行，返回行数。"""
    if not text:
        return 1
    lines, cur = 1, 0.0
    for ch in text:
        w = font.getlength(ch)
        if cur + w > width_pt and cur > 0:
            lines += 1
            cur = w
        else:
            cur += w
    return lines


def main(path):
    prs = Presentation(path)
    slide_w = prs.slide_width / 914400 * 72
    slide_h = prs.slide_height / 914400 * 72
    problems = []
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            x = sh.left / 914400 * 72
            y = sh.top / 914400 * 72
            w = sh.width / 914400 * 72
            h = sh.height / 914400 * 72
            if x < -1 or y < -1 or x + w > slide_w + 1 or y + h > slide_h + 1:
                problems.append((si, "越界", f"({x:.0f},{y:.0f},{w:.0f}x{h:.0f})",
                                 sh.text_frame.text[:30] if sh.has_text_frame else ""))
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            total = 0.0
            for p in sh.text_frame.paragraphs:
                runs = [r for r in p.runs if r.text]
                if not runs:
                    total += 6
                    continue
                size = runs[0].font.size.pt if runs[0].font.size else 12
                bold = bool(runs[0].font.bold)
                text = "".join(r.text for r in runs)
                f = face(runs[0].font.name or "PingFang SC", bold, size)
                n = wrap_lines(text, f, w)
                spacing = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
                total += n * size * spacing * 1.18
            if total > h * 1.06 + 1:
                problems.append((si, "溢出",
                                 f"需 {total:.0f}pt / 框 {h:.0f}pt（{w:.0f} 宽）",
                                 sh.text_frame.text.replace("\n", " / ")[:40]))
    for p in problems:
        print(f"P{p[0]:02d} [{p[1]}] {p[2]} :: {p[3]}")
    print(f"\n共 {len(problems)} 处")
    return problems


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "harness-tech-share-2026-10.pptx")
